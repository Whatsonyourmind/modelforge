"""DCF summary slide renderer -- assumptions table and sensitivity matrix with heatmap."""

from __future__ import annotations

import logging
import re
from typing import TYPE_CHECKING

from modelforge.deck.finance.conditional import ConditionalFormatter
from modelforge.deck.finance.formatter import FinancialFormatter
from modelforge.deck.ir.elements.base import Position
from modelforge.deck.rendering.slide_renderers.base import BaseFinanceSlideRenderer

if TYPE_CHECKING:
    from pptx.slide import Slide

    from modelforge.deck.ir.slides.base import BaseSlide
    from modelforge.deck.themes.types import ResolvedTheme

logger = logging.getLogger(__name__)

# Pattern to detect numeric-looking column headers (sensitivity grid axes).
_NUMERIC_HEADER = re.compile(r"^[\d.]+%?$")


class DcfSummaryRenderer(BaseFinanceSlideRenderer):
    """Renders DCF summary with assumptions table and color-coded sensitivity matrix."""

    def render(self, slide: Slide, ir_slide: BaseSlide, theme: ResolvedTheme) -> None:
        elements = ir_slide.elements

        # Title
        title = self._find_heading(elements) or "DCF Valuation Summary"
        title_pos = Position(x=0.75, y=0.4, width=11.8, height=0.8)
        self._add_title(slide, title, theme, title_pos)

        # Find table elements
        tables = self._find_table_elements(elements)
        if not tables:
            logger.warning("DcfSummaryRenderer: No table elements found")
            return

        # Classify tables: assumptions (key-value pairs) vs sensitivity (numeric grid)
        assumptions_table = None
        sensitivity_table = None

        for tbl in tables:
            content = tbl.content
            if self._is_sensitivity_table(content):
                sensitivity_table = tbl
            else:
                assumptions_table = tbl

        # Render assumptions table (left side)
        if assumptions_table:
            a_content = assumptions_table.content
            a_pos = Position(x=0.75, y=1.5, width=5.0, height=3.5)
            # Format assumption values
            formatted_rows = []
            for row in a_content.rows:
                formatted_row = list(row)
                if len(formatted_row) >= 2 and isinstance(formatted_row[1], (int, float)):
                    val = formatted_row[1]
                    # Heuristic: values < 1 are likely percentages
                    if isinstance(val, float) and 0 < val < 1:
                        formatted_row[1] = FinancialFormatter.percentage(val)
                    else:
                        formatted_row[1] = str(val)
                formatted_rows.append(formatted_row)

            self._add_table(
                slide,
                headers=a_content.headers,
                rows=formatted_rows,
                theme=theme,
                position=a_pos,
            )

        # Render sensitivity matrix (right side) with heatmap gradient
        if sensitivity_table:
            s_content = sensitivity_table.content
            s_pos = Position(x=6.5, y=1.5, width=6.0, height=3.5)

            # Compute min/max for gradient coloring
            all_values = []
            for row in s_content.rows:
                for val in row[1:]:  # Skip row header
                    if isinstance(val, (int, float)):
                        all_values.append(float(val))

            min_val = min(all_values) if all_values else 0
            max_val = max(all_values) if all_values else 1

            # Build cell_colors for gradient
            cell_colors: dict[tuple[int, int], str] = {}
            for row_idx, row in enumerate(s_content.rows):
                for col_idx in range(1, len(row)):
                    val = row[col_idx]
                    if isinstance(val, (int, float)):
                        color = ConditionalFormatter.heatmap_gradient(
                            float(val), min_val, max_val, theme
                        )
                        cell_colors[(row_idx, col_idx)] = color

            # Format as currency
            col_formats = [None] + ["currency"] * (len(s_content.headers) - 1)

            self._add_table(
                slide,
                headers=s_content.headers,
                rows=[list(r) for r in s_content.rows],
                theme=theme,
                position=s_pos,
                column_formats=col_formats,
                cell_colors=cell_colors,
            )

    @staticmethod
    def _is_sensitivity_table(content) -> bool:
        """Determine if a table content looks like a sensitivity matrix.

        Sensitivity tables have numeric-looking column headers (e.g., "1.5%", "2.0%").
        """
        if len(content.headers) < 3:
            return False
        # Check if most non-first headers look numeric
        remaining = content.headers[1:]
        numeric_count = sum(1 for h in remaining if _NUMERIC_HEADER.match(str(h)))
        return numeric_count >= len(remaining) * 0.5


__all__ = ["DcfSummaryRenderer"]
