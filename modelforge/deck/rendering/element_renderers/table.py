"""Table element renderer -- renders IR TableElement as python-pptx table shape."""

from __future__ import annotations

import logging
from typing import TYPE_CHECKING

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from modelforge.deck.rendering.element_renderers.base import BaseElementRenderer
from modelforge.deck.rendering.utils import hex_to_rgb, resolve_font_name

if TYPE_CHECKING:
    from pptx.slide import Slide

    from modelforge.deck.ir.elements.base import BaseElement, Position
    from modelforge.deck.themes.types import ResolvedTheme

logger = logging.getLogger(__name__)


def _is_numeric(value) -> bool:
    """Check if a value looks numeric (for alignment purposes)."""
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return True
    try:
        float(str(value).replace(",", "").replace("$", "").replace("%", ""))
        return True
    except (ValueError, TypeError):
        return False


class TableRenderer(BaseElementRenderer):
    """Renders table elements with header, body, optional footer, and highlight rows."""

    def render(self, slide: Slide, element: BaseElement, position: Position, theme: ResolvedTheme) -> None:
        content = element.content
        headers = content.headers
        rows = content.rows
        footer_row = content.footer_row
        highlight_rows = content.highlight_rows or []

        num_cols = len(headers)
        num_rows = 1 + len(rows)  # header + body rows
        if footer_row:
            num_rows += 1

        left = Inches(position.x)
        top = Inches(position.y)
        width = Inches(position.width)
        height = Inches(position.height)

        table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
        table = table_shape.table

        font_size = Pt(theme.typography.scale.get("caption", 14))
        font_family = resolve_font_name(theme.typography.body_family)

        # Distribute column widths proportionally based on header text lengths
        total_len = sum(max(len(str(h)), 3) for h in headers)
        for col_idx, header in enumerate(headers):
            proportion = max(len(str(header)), 3) / total_len
            table.columns[col_idx].width = int(width * proportion)

        # Header row
        for col_idx, header in enumerate(headers):
            cell = table.cell(0, col_idx)
            cell.text = str(header)

            # Style header
            self._style_cell(
                cell,
                font_size=font_size,
                font_family=font_family,
                bold=True,
                text_color=hex_to_rgb("#FFFFFF"),
                fill_color=hex_to_rgb(theme.colors.primary),
            )

        # Body rows
        for row_idx, row_data in enumerate(rows):
            table_row = row_idx + 1  # offset for header
            is_highlighted = row_idx in highlight_rows

            # Alternating row colors
            if is_highlighted:
                bg_color = hex_to_rgb(theme.colors.accent)
            elif row_idx % 2 == 0:
                bg_color = hex_to_rgb(theme.colors.surface)
            else:
                bg_color = hex_to_rgb(theme.colors.background)

            for col_idx in range(num_cols):
                cell_value = row_data[col_idx] if col_idx < len(row_data) else ""
                cell = table.cell(table_row, col_idx)
                cell.text = str(cell_value) if cell_value is not None else ""

                # Right-align numeric cells
                alignment = PP_ALIGN.RIGHT if _is_numeric(cell_value) else PP_ALIGN.LEFT

                self._style_cell(
                    cell,
                    font_size=font_size,
                    font_family=font_family,
                    bold=False,
                    text_color=hex_to_rgb(theme.colors.text_primary),
                    fill_color=bg_color,
                    alignment=alignment,
                )

        # Footer row
        if footer_row:
            footer_table_row = num_rows - 1
            for col_idx in range(num_cols):
                cell_value = footer_row[col_idx] if col_idx < len(footer_row) else ""
                cell = table.cell(footer_table_row, col_idx)
                cell.text = str(cell_value) if cell_value is not None else ""

                alignment = PP_ALIGN.RIGHT if _is_numeric(cell_value) else PP_ALIGN.LEFT

                self._style_cell(
                    cell,
                    font_size=font_size,
                    font_family=font_family,
                    bold=True,
                    text_color=hex_to_rgb(theme.colors.text_primary),
                    fill_color=hex_to_rgb(theme.colors.surface),
                    alignment=alignment,
                )

    def _style_cell(
        self,
        cell,
        font_size,
        font_family: str,
        bold: bool,
        text_color,
        fill_color,
        alignment=None,
    ) -> None:
        """Apply styling to a table cell."""
        # Cell fill
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill_color

        # Text styling
        for paragraph in cell.text_frame.paragraphs:
            if alignment:
                paragraph.alignment = alignment
            for run in paragraph.runs:
                run.font.size = font_size
                run.font.name = font_family
                run.font.bold = bold
                run.font.color.rgb = text_color


__all__ = ["TableRenderer"]
