"""What the moat gate counts, and what it must still catch.

Two refinements are locked here, each paired with the failure it must not stop
detecting — a gate that only ever passes is worth nothing.

1. ``formula_density`` excludes cells the workbook *declares* as inputs by
   painting them the blue input colour. Before, a core-output sheet carrying
   its own input block (an opening balance sheet, a period-index axis, a
   historical FY0 column) could not reach the threshold however live its
   outputs were. The check that matters is unaffected: ``auto_style_gaps``
   paints unstyled numerics black, so a pasted output is never blue.

2. ``no_orphan_inputs`` exempts the names the deck exporter reads. The gate
   asks whether anything consumes a defined name; for those the answer is yes,
   just not from a formula.
"""

from __future__ import annotations

import openpyxl
import pytest
from openpyxl.styles import Font

from modelforge.builder import styles
from modelforge.deck.named_facts import DECK_NAMED_FACTS
from modelforge.moat import MoatGate

BLUE = Font(color=styles.COLOR_INPUT)
BLACK = Font(color=styles.COLOR_FORMULA)


def _sheet(tmp_path, name, cells):
    """Build a one-sheet workbook; cells is [(coord, value, font)]."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = name
    for coord, value, font in cells:
        ws[coord] = value
        ws[coord].font = font
    path = tmp_path / f"{name.lower()}.xlsx"
    wb.save(path)
    return path


def _density(path):
    report = MoatGate().evaluate(path)
    return next(g for g in report.gate_results if g.name == "formula_density")


def test_declared_inputs_do_not_count_against_density(tmp_path):
    """Six blue inputs feeding four live formulas is a healthy sheet."""
    cells = [(f"B{i}", i * 1.5, BLUE) for i in range(1, 7)]
    cells += [(f"C{i}", f"=B{i}*2", BLACK) for i in range(1, 5)]
    gate = _density(_sheet(tmp_path, "Valuation", cells))
    assert gate.passed, gate.detail


def test_pasted_outputs_still_fail_density(tmp_path):
    """The IPO/restructuring defect: outputs pasted as values, styled black.

    This is the regression that matters. If excluding declared inputs ever
    starts excluding these too, the gate stops catching the thing it was
    built for.
    """
    cells = [(f"B{i}", i * 1.5, BLUE) for i in range(1, 4)]
    cells += [(f"C{i}", i * 3.0, BLACK) for i in range(1, 12)]  # should be formulas
    gate = _density(_sheet(tmp_path, "Valuation", cells))
    assert not gate.passed
    assert "Valuation" in gate.detail and "0.0%" in gate.detail


def test_unstyled_pasted_outputs_still_fail_density(tmp_path):
    """A cell with no explicit colour is not a declared input either."""
    cells = [(f"C{i}", i * 3.0, Font()) for i in range(1, 12)]
    assert not _density(_sheet(tmp_path, "Valuation", cells)).passed


def test_deck_named_facts_are_exempt_from_the_orphan_gate(tmp_path):
    """A name the deck reads is consumed, even though no formula names it."""
    assert "purchase_price_eur_m" in DECK_NAMED_FACTS
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Valuation"
    ws["D5"] = 250.0
    ws["D5"].font = BLUE
    ws["D6"] = "=D5*2"
    wb.defined_names.add(
        openpyxl.workbook.defined_name.DefinedName(
            "purchase_price_eur_m", attr_text="'Valuation'!$D$5")
    )
    path = tmp_path / "deck_anchor.xlsx"
    wb.save(path)
    report = MoatGate().evaluate(path)
    assert "purchase_price_eur_m" not in report.orphan_named_ranges


def test_a_genuinely_unread_name_is_still_an_orphan(tmp_path):
    """The exemption is a named list, not a blanket amnesty."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Valuation"
    ws["D5"] = 250.0
    ws["D5"].font = BLUE
    ws["D6"] = "=D5*2"
    wb.defined_names.add(
        openpyxl.workbook.defined_name.DefinedName(
            "a_knob_nothing_reads", attr_text="'Valuation'!$D$5")
    )
    path = tmp_path / "orphan.xlsx"
    wb.save(path)
    report = MoatGate().evaluate(path)
    assert "a_knob_nothing_reads" in report.orphan_named_ranges
